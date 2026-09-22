import docx
import pandas as pd
import json
import re

doc = docx.Document('data/Relatório de Pesquisa de Satisfação - FESUPER 2026 v2.docx')
with open('src/data/dashboard_data.json', 'r', encoding='utf-8') as f:
    dash = json.load(f)

df_v = pd.read_excel('data/FESUPER 2026_ PESQUISA DE SATISFAÇÃO  - VISITANTES.xlsx')
df_e = pd.read_excel('data/FESUPER 2026_ PESQUISA DE SATISFAÇÃO -  EXPOSITORES.xlsx')

print('====================================================================')
print('AUDITORIA DETALHADA: EXCEL BRUTO vs DOCX RELATÓRIO vs PAINEL DASHBOARD')
print('====================================================================')

# 1. AMOSTRA
print('\n--- 1. AMOSTRA TOTAL ---')
print(f'Expositores: Raw={len(df_e)}, Docx={doc.tables[0].rows[1].cells[1].text.strip()}, Dash={dash["metadata"]["total_expositores"]}')
print(f'Visitantes:  Raw={len(df_v)}, Docx={doc.tables[0].rows[1].cells[2].text.strip()}, Dash={dash["metadata"]["total_visitantes"]}')
print(f'Total:        Raw={len(df_e)+len(df_v)}, Docx={int(doc.tables[0].rows[1].cells[1].text.strip())+int(doc.tables[0].rows[1].cells[2].text.strip())}, Dash={dash["metadata"]["total_geral_respostas"]}')

# 2. EMPRESAS DISTINTAS
print('\n--- 2. EMPRESAS EXPOSITORAS DISTINTAS ---')
raw_companies = df_e.iloc[:, 6].dropna().str.strip()
print(f'Nomes brutos na coluna: {len(raw_companies)}')
print(f'Nomes unicos exatos: {raw_companies.nunique()}')
# Checar duplicatas
vc = raw_companies.value_counts()
dups = vc[vc > 1]
print('Empresas com mais de uma resposta no Excel:')
for comp, cnt in dups.items():
    print(f'  - {comp}: {cnt} respostas')
print(f'Docx Tabela 2 / Tabela 6 declara: 59 nomes distintos')

# 3. METODOLOGIA DE ESCALAS EM EXPOSITORES
print('\n--- 3. ESCALAS DE EXPOSITORES (DIMENSÕES OPERACIONAIS) ---')
# Colunas 8 a 16
dim_cols = [
    ('Organização geral', 8),
    ('Comunicação', 9),
    ('Credenciamento e recepção', 10),
    ('Estrutura e montagem', 11),
    ('Limpeza, conforto e segurança', 12),
    ('Apoio da organização', 13),
    ('Programação da feira', 14),
    ('Quantidade e perfil dos visitantes', 15),
    ('Experiência geral da empresa', 16)
]

midpoints = {
    '10': 10.0,
    '8-9': 8.5,
    '6-7': 6.5,
    '4-5': 4.5,
    '2-3': 2.5,
    '0-1': 0.5
}

# Pegar Tabela 7 do Docx
docx_t7 = {}
for r in doc.tables[6].rows[1:]:
    cells = [c.text.strip() for c in r.cells]
    dim_name = cells[0]
    valid = cells[1]
    media = cells[2]
    p810 = cells[3]
    docx_t7[dim_name] = {'valid': valid, 'media': media, 'p810': p810}

print(f"{'Dimensão':<35} | {'Raw Válidos':<11} | {'Docx Válidos':<12} | {'Raw Média':<10} | {'Docx Média':<10} | {'Dash Média':<10} | {'Status'}")
print('-'*105)

for name, col_idx in dim_cols:
    s = df_e.iloc[:, col_idx].dropna().astype(str).str.strip()
    valid_raw = len(s)
    
    # Calcular média estimada por ponto médio
    num_vals = []
    c_810 = 0
    for v in s:
        # mapear
        if v in midpoints:
            num_vals.append(midpoints[v])
            if v in ['10', '8-9']:
                c_810 += 1
        else:
            # tentar converter direto se for float
            try:
                fl = float(v)
                num_vals.append(fl)
                if fl >= 8.0:
                    c_810 += 1
            except:
                pass
    raw_mean = sum(num_vals)/len(num_vals) if num_vals else 0.0
    
    # Match com Docx
    docx_match = None
    for k in docx_t7:
        if k.lower()[:8] in name.lower():
            docx_match = docx_t7[k]
            break
            
    # Match com Dashboard
    dash_match = None
    for d in dash['dimensoes_operacionais']:
        if d['dimensao'].lower()[:8] in name.lower():
            dash_match = d
            break
            
    docx_val = docx_match['valid'] if docx_match else 'N/A'
    docx_m = docx_match['media'] if docx_match else 'N/A'
    dash_m = f"{dash_match['nota']:.2f}" if dash_match else 'N/A'
    
    status = 'OK' if (docx_m != 'N/A' and abs(raw_mean - float(docx_m.replace(',', '.'))) < 0.02 and dash_m == docx_m) else 'DIVERGÊNCIA'
    print(f"{name:<35} | {valid_raw:<11} | {docx_val:<12} | {raw_mean:<10.2f} | {docx_m:<10} | {dash_m:<10} | {status}")

# 4. VISITANTES: EXPERIÊNCIA, EXPECTATIVAS, RETORNO
print('\n--- 4. INDICADORES NUMÉRICOS DE VISITANTES ---')
# Col 11: Experiência, Col 12: Expectativas, Col 14: Retorno
vis_metrics = [
    ('Experiência geral', 11),
    ('Atendimento expectativas', 12),
    ('Intenção de retorno', 14)
]

# Docx Tabela 12
docx_t12 = {}
for r in doc.tables[11].rows[1:]:
    cells = [c.text.strip() for c in r.cells]
    docx_t12[cells[0]] = {'valid': cells[1], 'media': cells[2], 'p810': cells[3], 'p910': cells[4]}

for name, col_idx in vis_metrics:
    s = pd.to_numeric(df_v.iloc[:, col_idx], errors='coerce').dropna()
    valid_raw = len(s)
    mean_raw = s.mean()
    p810_raw = (s >= 8).mean() * 100
    p910_raw = (s >= 9).mean() * 100
    
    # Docx match
    d_m = None
    for k in docx_t12:
        if k.lower()[:6] in name.lower():
            d_m = docx_t12[k]
            break
            
    print(f"Indicador: {name}")
    print(f"  Raw:  Válidos={valid_raw}, Média={mean_raw:.2f}, 8-10={p810_raw:.1f}%, 9-10={p910_raw:.1f}%")
    if d_m:
        print(f"  Docx: Válidos={d_m['valid']}, Média={d_m['media']}, 8-10={d_m['p810']}, 9-10={d_m['p910']}")
    if name == 'Experiência geral':
        print(f"  Dash: Média={dash['kpis_principais']['visitantes']['media_experiencia']}")
    elif name == 'Atendimento expectativas':
        print(f"  Dash: Média={dash['kpis_principais']['visitantes']['media_expectativas']}")
    elif name == 'Intenção de retorno':
        print(f"  Dash: Média={dash['kpis_principais']['visitantes']['media_retorno']}, NPS={dash['kpis_principais']['visitantes']['nps']}")

# 5. NPS ADAPTADO
print('\n--- 5. AUDITORIA NPS ADAPTADO ---')
# Expositores Retorno: Coluna 21
s_ret_e = pd.to_numeric(df_e.iloc[:, 21], errors='coerce').dropna()
prom_e = (s_ret_e >= 9).sum()
neut_e = ((s_ret_e >= 7) & (s_ret_e <= 8)).sum()
detr_e = (s_ret_e <= 6).sum()
nps_e = ((prom_e - detr_e) / len(s_ret_e)) * 100

print(f"Expositores (Retorno - Col 21):")
print(f"  Raw: Válidos={len(s_ret_e)}, Promotores(9-10)={prom_e} ({prom_e/len(s_ret_e)*100:.1f}%), Neutros(7-8)={neut_e} ({neut_e/len(s_ret_e)*100:.1f}%), Detratores(0-6)={detr_e} ({detr_e/len(s_ret_e)*100:.1f}%)")
print(f"  Raw NPS: {nps_e:.1f} (arredondado: {round(nps_e)})")
print(f"  Docx Tabela 1: +66")
print(f"  Dash: {dash['kpis_principais']['expositores']['nps']} (Prom: {dash['kpis_principais']['expositores']['promotores_pct']}%, Neut: {dash['kpis_principais']['expositores']['neutros_pct']}%, Detr: {dash['kpis_principais']['expositores']['detratores_pct']}%)")

# Visitantes Retorno: Coluna 14
s_ret_v = pd.to_numeric(df_v.iloc[:, 14], errors='coerce').dropna()
prom_v = (s_ret_v >= 9).sum()
neut_v = ((s_ret_v >= 7) & (s_ret_v <= 8)).sum()
detr_v = (s_ret_v <= 6).sum()
nps_v = ((prom_v - detr_v) / len(s_ret_v)) * 100

print(f"\nVisitantes (Retorno - Col 14):")
print(f"  Raw: Válidos={len(s_ret_v)}, Promotores(9-10)={prom_v} ({prom_v/len(s_ret_v)*100:.1f}%), Neutros(7-8)={neut_v} ({neut_v/len(s_ret_v)*100:.1f}%), Detratores(0-6)={detr_v} ({detr_v/len(s_ret_v)*100:.1f}%)")
print(f"  Raw NPS: {nps_v:.1f} (arredondado: {round(nps_v)})")
print(f"  Docx Tabela 1: +92")
print(f"  Dash: {dash['kpis_principais']['visitantes']['nps']} (Prom: {dash['kpis_principais']['visitantes']['promotores_pct']}%, Neut: {dash['kpis_principais']['visitantes']['neutros_pct']}%, Detr: {dash['kpis_principais']['visitantes']['detratores_pct']}%)")

# 6. NEGÓCIOS GERADOS
print('\n--- 6. AUDITORIA: NEGÓCIOS GERADOS ---')
# Expositores Coluna 19
neg_e = df_e.iloc[:, 19].dropna().str.strip()
print('Expositores Negócios (Col 19):')
vc_ne = neg_e.value_counts()
for k, v in vc_ne.items():
    print(f"  Raw: '{k}': {v} ({v/len(neg_e)*100:.1f}%)")
print(f"Docx Tabela 8: Realizados=14 (22,6%), Oportunidades=42 (67,7%), Não avalia=6 (9,7%)")
print(f"Dash Expositores: Realizados={dash['kpis_principais']['expositores']['negocios_efetivos_pct']}%, Potenciais/Efetivos={dash['kpis_principais']['expositores']['negocios_efetivos_ou_potenciais_pct']}%")

# Visitantes Coluna 13
neg_v = df_v.iloc[:, 13].dropna().str.strip()
print('\nVisitantes Negócios (Col 13):')
vc_nv = neg_v.value_counts()
for k, v in vc_nv.items():
    print(f"  Raw: '{k}': {v} ({v/len(neg_v)*100:.1f}%)")
print(f"Docx Tabela 13: Já realizou=39 (32,0%), Pretende=33 (27,0%), Avaliando=26 (21,3%), Não=24 (19,7%)")
print(f"Dash Visitantes: Realizado={dash['kpis_principais']['visitantes']['negocio_realizado_pct']}%, Realizado ou Pretendido={dash['kpis_principais']['visitantes']['negocio_realizado_ou_pretendido_pct']}%")

# 7. REPRESENTANTES E EQUIPE
print('\n--- 7. AUDITORIA: EQUIPE / REPRESENTANTES DE EXPOSITORES ---')
# Coluna 22
s_col22 = df_e.iloc[:, 22].dropna().astype(str)
print(f"Respostas em Coluna 22: {len(s_col22)}")
print(f"Docx Tabela 6:")
for r in doc.tables[5].rows:
    print('  ' + ' | '.join([c.text.strip() for c in r.cells]))
print(f"Dash Expositores: Equipe média por estande = {dash['kpis_principais']['expositores']['equipe_media_por_estande']}, Total estimado = {dash['kpis_principais']['expositores']['total_profissionais_estimados']}")

# 8. AUDITORIA DA COBERTURA QUALITATIVA (FEEDBACKS)
print('\n--- 8. AUDITORIA: COBERTURA QUALITATIVA / FEEDBACKS ---')
# Expositores: Col 23 (Abaixo expectativa), Col 24 (Sugestão)
e_abaixo = df_e.iloc[:, 23].dropna().astype(str).str.strip()
e_sugest = df_e.iloc[:, 24].dropna().astype(str).str.strip()
e_ab_valid = [x for x in e_abaixo if x.lower() not in ['não', 'nao', 'nenhum', 'nenhuma', 'não.', 'nao.', 'n']]
e_sug_valid = [x for x in e_sugest if x.lower() not in ['não', 'nao', 'nenhuma', 'sem sugestões', 'nada', 'não.', 'nao.']]

# Visitantes: Col 15 (Abaixo expectativa), Col 16 (Sugestão)
v_abaixo = df_v.iloc[:, 15].dropna().astype(str).str.strip()
v_sugest = df_v.iloc[:, 16].dropna().astype(str).str.strip()
v_ab_valid = [x for x in v_abaixo if x.lower() not in ['não', 'nao', 'nenhum', 'nenhuma', 'não.', 'nao.', 'n', 'tudo ok']]
v_sug_valid = [x for x in v_sugest if x.lower() not in ['não', 'nao', 'nenhuma', 'sem sugestões', 'nada', 'não.', 'nao.', 'nenhuma sugestão']]

print(f"Expositores: Col 23 respostas com texto={len(e_abaixo)} (válidos substantivos={len(e_ab_valid)}), Col 24 respostas com texto={len(e_sugest)} (válidos={len(e_sug_valid)})")
print(f"Visitantes:  Col 15 respostas com texto={len(v_abaixo)} (válidos substantivos={len(v_ab_valid)}), Col 16 respostas com texto={len(v_sugest)} (válidos={len(v_sug_valid)})")
print(f"Total feedbacks catalogados no dashboard: {len(dash['feedbacks'])}")

dash_texts = [f['texto'].lower() for f in dash['feedbacks']]

print('\n--- 9. AUDITORIA DE CITAÇÕES LITERAIS DO DOCX NO PAINEL ---')
print('Expositores (Docx Tabela 9):')
t9 = doc.tables[8]
for r in t9.rows[1:]:
    tema = r.cells[0].text.strip()
    sug = r.cells[1].text.strip().replace('“', '').replace('”', '').replace('"', '')
    clean_search = sug.lower()[:20]
    found = any(clean_search in dt for dt in dash_texts)
    print(f"  [{tema:<22}] '{sug}' -> Presente no painel: {found}")

print('\nVisitantes (Docx Tabela 14):')
t14 = doc.tables[13]
for r in t14.rows[1:]:
    tema = r.cells[0].text.strip()
    sug = r.cells[1].text.strip().replace('“', '').replace('”', '').replace('"', '')
    clean_search = sug.lower()[:20]
    found = any(clean_search in dt for dt in dash_texts)
    print(f"  [{tema:<22}] '{sug}' -> Presente no painel: {found}")
